"""Builds the DOCX versions of the Markdown docs (ТЗ, section 5: README + DOCX).

Run from the project root:  python tool/build_docs.py
Needs python-docx. Converts headings, paragraphs, lists, tables, quotes and
code blocks; the RuStore card also gets its icon and screenshots.
"""
import os
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCS = os.path.join(ROOT, 'docs')

JOBS = [
    ('DOCUMENTATION.md', 'Документация_Финансовый_питомец.docx', None),
    ('rustore/RUSTORE_CARD.md', 'rustore/Карточка_RuStore_Финансовый_питомец.docx', 'rustore'),
]

INLINE = re.compile(r'(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))')
NAVY = RGBColor(0x08, 0x15, 0x4B)


def add_inline(paragraph, text):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
        elif part.startswith('['):
            label = part[1:part.index(']')]
            paragraph.add_run(label.replace('`', ''))
        else:
            paragraph.add_run(part)


def shade(cell, color):
    properties = cell._tc.get_or_add_tcPr()
    fill = OxmlElement('w:shd')
    fill.set(qn('w:val'), 'clear')
    fill.set(qn('w:color'), 'auto')
    fill.set(qn('w:fill'), color)
    properties.append(fill)


def add_table(document, rows):
    cells = [[c.strip() for c in row.strip().strip('|').split('|')] for row in rows]
    header, body = cells[0], [r for r in cells[2:]]
    table = document.add_table(rows=1, cols=len(header))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, text in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.paragraphs[0].text = ''
        add_inline(cell.paragraphs[0], text)
        for run in cell.paragraphs[0].runs:
            run.bold = True
        shade(cell, 'DCEEFF')
    for row in body:
        new = table.add_row().cells
        for i in range(len(header)):
            add_inline(new[i].paragraphs[0], row[i] if i < len(row) else '')
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9.5)
    document.add_paragraph()


def convert(markdown, document):
    lines = markdown.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if stripped.startswith('```'):
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                block.append(lines[i])
                i += 1
            paragraph = document.add_paragraph()
            run = paragraph.add_run('\n'.join(block))
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
            paragraph.paragraph_format.left_indent = Cm(0.5)
        elif stripped.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                rows.append(lines[i])
                i += 1
            add_table(document, rows)
            continue
        elif stripped.startswith('#'):
            level = len(stripped) - len(stripped.lstrip('#'))
            heading = document.add_heading(level=min(level, 3) if level > 1 else 0)
            add_inline(heading, stripped.lstrip('#').strip())
            for run in heading.runs:
                run.font.color.rgb = NAVY
        elif stripped == '---':
            pass
        elif stripped.startswith('>'):
            text = stripped.lstrip('>').strip()
            if text:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Cm(0.8)
                add_inline(paragraph, text)
        elif re.match(r'^(\s*)[-•] ', line):
            indent = len(line) - len(line.lstrip())
            paragraph = document.add_paragraph(
                style='List Bullet 2' if indent >= 2 else 'List Bullet')
            add_inline(paragraph, re.sub(r'^\s*[-•] ', '', line))
        elif re.match(r'^\d+\. ', stripped):
            paragraph = document.add_paragraph(style='List Number')
            add_inline(paragraph, re.sub(r'^\d+\. ', '', stripped))
        elif stripped:
            paragraph = document.add_paragraph()
            add_inline(paragraph, stripped)
        i += 1


def add_store_images(document):
    document.add_heading('Иконка и скриншоты', level=2)
    document.add_picture(os.path.join(DOCS, 'rustore/icon_512.png'), width=Cm(4))
    table = document.add_table(rows=2, cols=3)
    for index in range(6):
        cell = table.rows[index // 3].cells[index % 3]
        cell.paragraphs[0].add_run().add_picture(
            os.path.join(DOCS, f'rustore/store/screenshot_{index + 1}.png'), width=Cm(5))


def main():
    for source, target, extras in JOBS:
        document = Document()
        style = document.styles['Normal']
        style.font.name = 'Calibri'
        style.font.size = Pt(11)
        for section in document.sections:
            section.left_margin = section.right_margin = Cm(2)
        with open(os.path.join(DOCS, source), encoding='utf-8') as file:
            convert(file.read(), document)
        if extras == 'rustore':
            add_store_images(document)
        document.save(os.path.join(DOCS, target))
        print('written', target)


if __name__ == '__main__':
    main()
